"""Respaldo Excel de 9 hojas, en español, sin campos base64."""
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from db import (
    users,
    clients,
    credit_processes,
    technical_processes,
    construction_schedules,
    payment_schedules,
    audit_logs,
)
from labels import (
    ROLES,
    USER_STATUS,
    CLIENT_STATUS,
    PAYMENT_METHOD,
    STAGE_STATUS,
    INSTALLMENT_STATUS,
    FREQUENCY,
    CREDIT_STAGES,
    TECHNICAL_STAGES,
    CREDIT_STAGE_LABELS,
    TECHNICAL_STAGE_LABELS,
)
from utils import compute_installment_status, parse_iso_date

# Colores corporativos DOMO.
AZUL = "1A3A6B"
NARANJA = "F57C00"
VERDE = "4CAF50"
ROJO = "D32F2F"
NARANJA_ALERTA = "FF9800"
GRIS = "ECEFF1"

HEADER_FILL = PatternFill("solid", fgColor=AZUL)
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
SECTION_FONT = Font(color="FFFFFF", bold=True, size=12)
THIN = Side(style="thin", color="CFD8DC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

STATUS_EMOJI = {
    "pending": "🟠",
    "in_progress": "🟡",
    "completed": "🟢",
    "paid": "🟢",
    "overdue": "🔴",
}


def _fmt_dt(value):
    if not value:
        return ""
    if isinstance(value, str):
        return value[:19].replace("T", " ")
    try:
        return value.strftime("%Y-%m-%d %H:%M")
    except Exception:
        return str(value)


def _fmt_date(value):
    if not value:
        return ""
    if isinstance(value, str):
        return value[:10]
    try:
        return value.strftime("%Y-%m-%d")
    except Exception:
        return str(value)


def _write_headers(ws, headers, row=1):
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def _autosize(ws, max_width=55):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            # Ignorar celdas combinadas (no exponen column_letter de forma segura).
            if cell.__class__.__name__ == "MergedCell":
                continue
            val = "" if cell.value is None else str(cell.value)
            idx = cell.column
            widths[idx] = max(widths.get(idx, 0), len(val))
    for idx, length in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = min(
            max_width, max(12, length + 2)
        )


async def build_excel() -> bytes:
    wb = Workbook()

    users_docs = await users.find().sort("created_at", -1).to_list(5000)
    clients_docs = await clients.find().sort("created_at", -1).to_list(5000)
    credit_docs = {
        d["client_id"]: d for d in await credit_processes.find().to_list(5000)
    }
    tech_docs = {
        d["client_id"]: d for d in await technical_processes.find().to_list(5000)
    }
    constr_docs = {
        d["client_id"]: d for d in await construction_schedules.find().to_list(5000)
    }
    pay_docs = {
        d["client_id"]: d for d in await payment_schedules.find().to_list(5000)
    }
    audit_docs = (
        await audit_logs.find().sort("created_at", -1).to_list(2000)
    )

    # ---------- Hoja 1: Resumen ----------
    ws = wb.active
    ws.title = "Resumen"
    ws.merge_cells("A1:C1")
    t = ws["A1"]
    t.value = "DOMO Constructora — Respaldo de datos"
    t.font = Font(color="FFFFFF", bold=True, size=14)
    t.fill = HEADER_FILL
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    total_cobrado = total_pendiente = total_vencido = 0.0
    for p in pay_docs.values():
        for inst in p.get("installments", []):
            st = compute_installment_status(inst)
            amt = inst.get("amount", 0) or 0
            if st == "paid":
                total_cobrado += amt
            elif st == "overdue":
                total_vencido += amt
                total_pendiente += amt
            else:
                total_pendiente += amt

    rows = [
        ("Fecha de generación", date.today().isoformat()),
        ("Total de usuarios", len(users_docs)),
        ("Total de clientes", len(clients_docs)),
        ("Clientes por crédito", sum(1 for c in clients_docs if c.get("payment_method") == "credito")),
        ("Clientes por efectivo", sum(1 for c in clients_docs if c.get("payment_method") == "efectivo")),
        ("Total cobrado (USD)", round(total_cobrado, 2)),
        ("Total pendiente (USD)", round(total_pendiente, 2)),
        ("Total vencido (USD)", round(total_vencido, 2)),
    ]
    r = 3
    for label, value in rows:
        c1 = ws.cell(row=r, column=1, value=label)
        c1.font = Font(bold=True)
        c1.fill = PatternFill("solid", fgColor=GRIS)
        ws.cell(row=r, column=2, value=value)
        r += 1
    _autosize(ws)

    # ---------- Hoja 2: Clientes ----------
    ws = wb.create_sheet("Clientes")
    headers = ["Nombre", "Teléfono", "Correo", "Dirección", "Método de pago",
               "Estado", "Notas", "Creado por", "Fecha creación"]
    _write_headers(ws, headers)
    for i, c in enumerate(clients_docs, start=2):
        ws.cell(row=i, column=1, value=c.get("name"))
        ws.cell(row=i, column=2, value=c.get("phone"))
        ws.cell(row=i, column=3, value=c.get("email"))
        ws.cell(row=i, column=4, value=c.get("address"))
        ws.cell(row=i, column=5, value=PAYMENT_METHOD.get(c.get("payment_method"), ""))
        ws.cell(row=i, column=6, value=CLIENT_STATUS.get(c.get("status"), ""))
        ws.cell(row=i, column=7, value=c.get("notes"))
        ws.cell(row=i, column=8, value=c.get("created_by_name"))
        ws.cell(row=i, column=9, value=_fmt_dt(c.get("created_at")))
    _autosize(ws)

    # ---------- Hoja 3: Detalle por Cliente ----------
    ws = wb.create_sheet("Detalle por Cliente")
    row = 1
    for c in clients_docs:
        cid = str(c["_id"])
        # Encabezado del cliente.
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        cell = ws.cell(row=row, column=1, value=f"👤 {c.get('name')}  ·  {PAYMENT_METHOD.get(c.get('payment_method'),'')}  ·  {CLIENT_STATUS.get(c.get('status'),'')}")
        cell.fill = HEADER_FILL
        cell.font = SECTION_FONT
        row += 1
        ws.cell(row=row, column=1, value=f"Tel: {c.get('phone','')}  |  Correo: {c.get('email','')}")
        row += 2

        # Crédito
        if c.get("payment_method") != "efectivo":
            _detail_section(ws, row, "💳 Ruta de Crédito", NARANJA)
            row += 1
            cp = credit_docs.get(cid, {}).get("stages", {})
            for key in CREDIT_STAGES:
                st = cp.get(key, {})
                emoji = STATUS_EMOJI.get(st.get("status"), "⚪")
                ws.cell(row=row, column=1, value=f"  {emoji} {CREDIT_STAGE_LABELS.get(key, key)}")
                ws.cell(row=row, column=2, value=STAGE_STATUS.get(st.get("status"), "Pendiente"))
                ws.cell(row=row, column=3, value=st.get("notes", ""))
                ws.cell(row=row, column=4, value=_fmt_dt(st.get("updated_at")))
                row += 1
            row += 1

        # Técnico
        _detail_section(ws, row, "📐 Ruta Técnica", AZUL)
        row += 1
        tp = tech_docs.get(cid, {}).get("stages", {})
        for key in TECHNICAL_STAGES:
            st = tp.get(key, {})
            emoji = STATUS_EMOJI.get(st.get("status"), "⚪")
            ws.cell(row=row, column=1, value=f"  {emoji} {TECHNICAL_STAGE_LABELS.get(key, key)}")
            ws.cell(row=row, column=2, value=STAGE_STATUS.get(st.get("status"), "Pendiente"))
            ws.cell(row=row, column=3, value=st.get("notes", ""))
            ws.cell(row=row, column=4, value=_fmt_dt(st.get("updated_at")))
            row += 1
        row += 1

        # Obra
        _detail_section(ws, row, "🏗️ Cronograma de Obra", VERDE)
        row += 1
        acts = constr_docs.get(cid, {}).get("activities", [])
        if not acts:
            ws.cell(row=row, column=1, value="  (sin actividades)")
            row += 1
        for a in acts:
            emoji = STATUS_EMOJI.get(a.get("status"), "⚪")
            ws.cell(row=row, column=1, value=f"  {emoji} {a.get('name')}")
            ws.cell(row=row, column=2, value=STAGE_STATUS.get(a.get("status"), "Pendiente"))
            ws.cell(row=row, column=3, value=f"{a.get('progress', 0)}%")
            ws.cell(row=row, column=4, value=f"{_fmt_date(a.get('start_date'))} → {_fmt_date(a.get('end_date'))}")
            row += 1
        row += 1

        # Pagos
        _detail_section(ws, row, "💵 Plan de Pagos", NARANJA_ALERTA)
        row += 1
        pay = pay_docs.get(cid)
        if pay:
            ws.cell(row=row, column=1, value=f"  Total: ${pay.get('total_amount',0)}  |  Abono: ${pay.get('down_payment',0)}  |  Cuotas: {pay.get('num_installments',0)}")
            row += 1
            for inst in pay.get("installments", []):
                st = compute_installment_status(inst)
                emoji = STATUS_EMOJI.get(st, "⚪")
                ws.cell(row=row, column=1, value=f"  {emoji} Cuota #{inst.get('number')}")
                ws.cell(row=row, column=2, value=f"${inst.get('amount',0)}")
                ws.cell(row=row, column=3, value=_fmt_date(inst.get("due_date")))
                ws.cell(row=row, column=4, value=INSTALLMENT_STATUS.get(st, ""))
                row += 1
        else:
            ws.cell(row=row, column=1, value="  (sin plan de pagos)")
            row += 1
        row += 2  # separación entre clientes
    _autosize(ws, max_width=60)

    # ---------- Hoja 4: Ruta Crédito ----------
    ws = wb.create_sheet("Ruta Crédito")
    headers = ["Cliente", "Etapa", "Estado", "Notas", "Actualizado por", "Fecha"]
    _write_headers(ws, headers)
    r = 2
    for c in clients_docs:
        cid = str(c["_id"])
        stages = credit_docs.get(cid, {}).get("stages", {})
        for key in CREDIT_STAGES:
            st = stages.get(key, {})
            ws.cell(row=r, column=1, value=c.get("name"))
            ws.cell(row=r, column=2, value=CREDIT_STAGE_LABELS.get(key, key))
            ws.cell(row=r, column=3, value=STAGE_STATUS.get(st.get("status"), "Pendiente"))
            ws.cell(row=r, column=4, value=st.get("notes", ""))
            ws.cell(row=r, column=5, value=st.get("updated_by_name", ""))
            ws.cell(row=r, column=6, value=_fmt_dt(st.get("updated_at")))
            r += 1
    _autosize(ws)

    # ---------- Hoja 5: Ruta Técnica ----------
    ws = wb.create_sheet("Ruta Técnica")
    _write_headers(ws, headers)
    r = 2
    for c in clients_docs:
        cid = str(c["_id"])
        stages = tech_docs.get(cid, {}).get("stages", {})
        for key in TECHNICAL_STAGES:
            st = stages.get(key, {})
            ws.cell(row=r, column=1, value=c.get("name"))
            ws.cell(row=r, column=2, value=TECHNICAL_STAGE_LABELS.get(key, key))
            ws.cell(row=r, column=3, value=STAGE_STATUS.get(st.get("status"), "Pendiente"))
            ws.cell(row=r, column=4, value=st.get("notes", ""))
            ws.cell(row=r, column=5, value=st.get("updated_by_name", ""))
            ws.cell(row=r, column=6, value=_fmt_dt(st.get("updated_at")))
            r += 1
    _autosize(ws)

    # ---------- Hoja 6: Cronograma Obra ----------
    ws = wb.create_sheet("Cronograma Obra")
    headers = ["Cliente", "Actividad", "Estado", "Progreso", "Inicio", "Fin", "Notas", "Actualizado por"]
    _write_headers(ws, headers)
    r = 2
    for c in clients_docs:
        cid = str(c["_id"])
        acts = constr_docs.get(cid, {}).get("activities", [])
        for a in acts:
            ws.cell(row=r, column=1, value=c.get("name"))
            ws.cell(row=r, column=2, value=a.get("name"))
            ws.cell(row=r, column=3, value=STAGE_STATUS.get(a.get("status"), "Pendiente"))
            ws.cell(row=r, column=4, value=f"{a.get('progress', 0)}%")
            ws.cell(row=r, column=5, value=_fmt_date(a.get("start_date")))
            ws.cell(row=r, column=6, value=_fmt_date(a.get("end_date")))
            ws.cell(row=r, column=7, value=a.get("notes", ""))
            ws.cell(row=r, column=8, value=a.get("updated_by_name", ""))
            r += 1
    _autosize(ws)

    # ---------- Hoja 7: Plan de Pagos ----------
    ws = wb.create_sheet("Plan de Pagos")
    headers = ["Cliente", "Cuota #", "Monto", "Vencimiento", "Estado", "Fecha pago", "Notas"]
    _write_headers(ws, headers)
    r = 2
    for c in clients_docs:
        cid = str(c["_id"])
        pay = pay_docs.get(cid)
        if not pay:
            continue
        for inst in pay.get("installments", []):
            st = compute_installment_status(inst)
            ws.cell(row=r, column=1, value=c.get("name"))
            ws.cell(row=r, column=2, value=inst.get("number"))
            ws.cell(row=r, column=3, value=inst.get("amount"))
            ws.cell(row=r, column=4, value=_fmt_date(inst.get("due_date")))
            ws.cell(row=r, column=5, value=INSTALLMENT_STATUS.get(st, ""))
            ws.cell(row=r, column=6, value=_fmt_date(inst.get("paid_date")))
            ws.cell(row=r, column=7, value=inst.get("notes", ""))
            r += 1
    _autosize(ws)

    # ---------- Hoja 8: Usuarios ----------
    ws = wb.create_sheet("Usuarios")
    headers = ["Nombre", "Correo", "Rol", "Estado", "Fecha creación"]
    _write_headers(ws, headers)
    for i, u in enumerate(users_docs, start=2):
        ws.cell(row=i, column=1, value=u.get("name"))
        ws.cell(row=i, column=2, value=u.get("email"))
        ws.cell(row=i, column=3, value=ROLES.get(u.get("role"), u.get("role")))
        ws.cell(row=i, column=4, value=USER_STATUS.get(u.get("status"), u.get("status")))
        ws.cell(row=i, column=5, value=_fmt_dt(u.get("created_at")))
    _autosize(ws)

    # ---------- Hoja 9: Auditoría ----------
    ws = wb.create_sheet("Auditoría")
    headers = ["Fecha", "Usuario", "Rol", "Proceso", "Acción", "Descripción"]
    _write_headers(ws, headers)
    for i, a in enumerate(audit_docs, start=2):
        ws.cell(row=i, column=1, value=_fmt_dt(a.get("created_at")))
        ws.cell(row=i, column=2, value=a.get("user_name"))
        ws.cell(row=i, column=3, value=ROLES.get(a.get("user_role"), a.get("user_role")))
        ws.cell(row=i, column=4, value=a.get("process_type"))
        ws.cell(row=i, column=5, value=a.get("action"))
        ws.cell(row=i, column=6, value=a.get("description"))
    _autosize(ws)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _detail_section(ws, row, title, color):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    cell = ws.cell(row=row, column=1, value=title)
    cell.fill = PatternFill("solid", fgColor=color)
    cell.font = Font(color="FFFFFF", bold=True)
